#!/usr/bin/env python3
"""Generate small, public-domain CI test documents. No user files or paid services."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent / "generated"
ROOT.mkdir(exist_ok=True)
C='<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
OD='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
X='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def save(name, entries):
    with ZipFile(ROOT/name,'w',compression=ZIP_DEFLATED) as z:
        for key,val in entries.items():z.writestr(key,val)
rootrel=lambda target:C+f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{OD}/officeDocument" Target="{target}"/></Relationships>'
default_types=C+'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
word_text='OPENDESK OFFICE TEST 2026'
worddoc=C+f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>{word_text}</w:t></w:r></w:p><w:sectPr/></w:body></w:document>'
save('real_office.docx',{
 '[Content_Types].xml':default_types+'<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
 '_rels/.rels':rootrel('word/document.xml'),
 'word/document.xml':worddoc
})
sheets=C+f'<workbook xmlns="{X}" xmlns:r="{OD}"><sheets><sheet name="OpenDesk" sheetId="1" r:id="rId1"/></sheets></workbook>'
worksheet=C+f'<worksheet xmlns="{X}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>OPENDESK SHEET 2026</t></is></c></row></sheetData></worksheet>'
save('real_sheet.xlsx',{
 '[Content_Types].xml':default_types+'<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>',
 '_rels/.rels':rootrel('xl/workbook.xml'),
 'xl/workbook.xml':sheets,
 'xl/_rels/workbook.xml.rels':C+f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{OD}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
 'xl/worksheets/sheet1.xml':worksheet
})
im=Image.new('RGB',(1500,330),'white')
draw=ImageDraw.Draw(im)
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',84)
draw.text((75,75),'OPENDESK 2026',font=font,fill='black')
im.save(ROOT/'real_scanned_english.png',dpi=(300,300))

# Real, two-slide OOXML PowerPoint with vector shapes, text and an embedded image.
# Only generated public CI content is included.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
deck=Presentation()
deck.slide_width=Inches(13.333)
deck.slide_height=Inches(7.5)
for heading in ('OPENDESK SLIDE ONE 2026', 'OPENDESK SLIDE TWO 2026'):
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    rect=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.5), Inches(.5), Inches(12.3), Inches(.8))
    rect.fill.solid()
    rect.fill.fore_color.rgb=RGBColor(15,74,109)
    rect.line.fill.background()
    title=slide.shapes.add_textbox(Inches(1), Inches(1.7), Inches(11.3), Inches(1))
    title.text_frame.text=heading
    title.text_frame.paragraphs[0].runs[0].font.size=Pt(36)
    slide.shapes.add_picture(str(ROOT/'real_scanned_english.png'), Inches(1), Inches(3), width=Inches(7.5))
deck.save(ROOT/'real_slides.pptx')

print('Generated public test documents:',*[p.name for p in sorted(ROOT.iterdir())])
